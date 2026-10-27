"""The build workbook of one harness: a spreadsheet to build the harness from (decision D-169).

Where the other tables are data, this file is a document: a first sheet that says what is ready and
what is still missing, a cut list in the order a wire is cut, one build block per connector, the
parts, the accessories (sleeving, heat shrink, labels, backshells), the labels and a test sheet with
columns for the result. It holds only what the model holds; an empty cell is shown as *to fill in*
and nothing is invented. Strip lengths, contact parts and crimp tools (decision D-170) are always
columns here, because an empty cell is a to-do for the person who fills in the design. All values are worked out here, the
workbook has no formulas, and the same model gives the same bytes (`exports.normalise_xlsx`).
"""

import contextlib
import io
from collections import defaultdict
from dataclasses import dataclass
from typing import Any

from harness_design_studio.core.generate.lengths import wire_length
from harness_design_studio.core.model import Connector, Harness, Project, Wire
from harness_design_studio.core.wirecolours import resolve

from . import tables
from .exports import _FIXED, normalise_xlsx
from .stamp import Stamp, Table
from .xlsx_env import ILLEGAL_CHARACTERS_RE, Workbook, WriteOnlyCell

FILE_SUFFIX = "_build.xlsx"

SHEETS = (
    "Start here", "Cut list", "Connectors", "Parts", "Accessories", "Labels", "Tests", "About"
)  # fmt: skip

CUT_HEADER = [
    "Cut (x)", "Wire", "Colour", "AWG", "Part", "Length (m)", "Strip A (mm)", "Strip B (mm)", "From",
    "To", "Signal or label", "Shield group", "Note",
]  # fmt: skip
CONNECTOR_HEADER = [
    "Pin", "Signal", "Function", "Wire", "Colour", "AWG", "Goes to", "Termination", "Contact size",
    "Contact", "Tool", "Terminated (x)", "Inserted (x)", "Checked (x)",
]  # fmt: skip
ACCESSORY_HEADER = ["Done (x)", *tables.ACCESSORY_HEADER]
NO_ACCESSORIES = "No accessories entered."
HARNESS_TARGET = "(harness)"  # an accessory with no target belongs to the harness as a whole
LABEL_HEADER = ["Done (x)", "Label", "Kind", "Text", "Quantity"]
TEST_HEADER = [
    "Test", "Type", "From", "To", "Expected", "Limit", "Test voltage", "Wire", "Measured",
    "Result", "Initials", "Date",
]  # fmt: skip
PARTS_HEADER = ["Part", "Category", "Manufacturer", "Part number", "Description", "Approval",
                "Quantity", "Unit", "Wires without length", "Used in"]  # fmt: skip
CHECK_HEADER = ["Check", "Open", "Status", "What to do"]

TO_FILL = "to fill in"
TO_FIX = "TO FIX"
NOTE = "NOTE"
OK = "OK"

# colours of the layout (the wire colours come from core/wirecolours.py)
_NAVY, _WHITE, _INK = "1F3A5F", "FFFFFF", "1A1A1A"
_BAND, _GREY_TEXT, _SPARE = "E8EEF5", "6B6B6B", "F0F0F0"
_RED, _AMBER, _GREEN = "F8CBAD", "FFE699", "C6EFCE"


@dataclass
class _Layout:
    """Column widths seen while a sheet is written (the widest text, capped)."""

    widths: dict[int, float]

    def see(self, col: int, text: str, cap: float = 46.0) -> None:
        longest = max((len(part) for part in text.split("\n")), default=0)
        self.widths[col] = min(cap, max(self.widths.get(col, 0.0), longest + 2.0))


def _styles() -> dict[str, Any]:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    thin = Side(style="thin", color="B7B7B7")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)

    def fill(rgb: str) -> PatternFill:
        return PatternFill("solid", start_color=rgb, end_color=rgb)

    return {
        "Font": Font, "Alignment": Alignment, "Border": Border, "Side": Side, "fill": fill,
        "box": box, "wrap": Alignment(wrap_text=True, vertical="top"),
        "centre": Alignment(horizontal="center", vertical="top"),
        "flat": Alignment(wrap_text=False, vertical="top"),
    }  # fmt: skip


_S = _styles()


def _text(value: object) -> str:
    return str(ILLEGAL_CHARACTERS_RE.sub("", str(value)))


def _cell(
    ws: Any,
    value: object,
    *,
    bold: bool = False,
    fill: str | None = None,
    colour: str | None = None,
    size: float | None = None,
    box: bool = True,
    centre: bool = False,
    italic: bool = False,
    wrap: bool = True,
) -> Any:
    """One styled cell. Text that a spreadsheet program would run as a formula (= + - @) is stored
    as text; numbers stay numbers."""
    c = WriteOnlyCell(ws, value=value if isinstance(value, int | float) else _text(value))
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        c.data_type = "s"
    c.font = _S["Font"](bold=bold, italic=italic, color=colour or _INK, size=size or 10)
    if fill:
        c.fill = _S["fill"](fill)
    if box:
        c.border = _S["box"]
    c.alignment = _S["centre"] if centre else _S["wrap"] if wrap else _S["flat"]
    return c


def _contrast(rgb: str) -> str:
    """Black or white text, whichever reads better on this fill."""
    r, g, b = (int(rgb[i : i + 2], 16) for i in (1, 3, 5))
    return "FFFFFF" if 0.299 * r + 0.587 * g + 0.114 * b < 140 else _INK


def _wire_colour_cell(ws: Any, name: str | None) -> Any:
    """The colour as a coloured cell with its name beside it in words (never colour alone)."""
    if not name:
        return _cell(ws, TO_FILL, fill=_AMBER)
    code, screen = resolve(name)
    return _cell(ws, _colour_text(name), fill=screen[1:], colour=_contrast(screen))


def _colour_text(name: str | None) -> str:
    """The colour in words, with its IEC 60757 code when the table knows it: "red (RD)"."""
    if not name:
        return ""
    code, _screen = resolve(name)
    return f"{name} ({code})" if code and code != name.upper() else name


def _header(ws: Any, lay: _Layout, labels: list[str]) -> list[Any]:
    out = []
    for n, label in enumerate(labels, start=1):
        lay.see(n, label)
        out.append(_cell(ws, label, bold=True, fill=_NAVY, colour=_WHITE))
    return out


def _finish(
    ws: Any,
    lay: _Layout,
    stamp: Stamp,
    body: list[list[Any]],
    *,
    tab: str,
    header_row: int | None = 1,
    freeze: str | None = None,
    one_page: bool = False,
) -> None:
    """Widths, print set-up and the footer every sheet gets, then the rows. In a write-only
    workbook all of this must be set before the first row is written, so a sheet collects its
    rows first."""
    if freeze:
        ws.freeze_panes = freeze
    for col, width in lay.widths.items():
        ws.column_dimensions[_letter(col)].width = round(width, 1)
    ws.sheet_properties.tabColor = tab
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = 9  # A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1 if one_page else 0
    from openpyxl.worksheet.properties import PageSetupProperties

    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = 0.6
    ws.page_margins.bottom = 0.8
    ws.page_margins.footer = 0.3
    if header_row:
        ws.print_title_rows = f"{header_row}:{header_row}"
    ws.oddFooter.left.text = stamp.line
    ws.oddFooter.center.text = "&A"
    ws.oddFooter.right.text = "Page &P of &N"
    for cells in body:
        ws.append(cells)


def _letter(col: int) -> str:
    from openpyxl.utils import get_column_letter

    return str(get_column_letter(col))


def _tick_rules(ws: Any, ref: str) -> None:
    """A cell that holds an x turns green; the list offers the x so nobody has to type it."""
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.worksheet.datavalidation import DataValidation

    dv = DataValidation(type="list", formula1='"x"', allow_blank=True, showDropDown=False)
    ws.data_validations.append(dv)
    dv.add(ref)
    first = ref.split(":")[0]
    ws.conditional_formatting.add(
        ref,
        FormulaRule(formula=[f"LEN({first})>0"], fill=_S["fill"](_GREEN), stopIfTrue=False),
    )


# ---- what each sheet needs from the model ----------------------------------------------------


def _wire_sort_key(h: Harness, w: Wire) -> tuple[int, int, str, int, float, str]:
    """Cut order: thick before thin, then by colour, then by length, so the same wire is cut in
    one go; wires with no gauge or no length come last."""
    length = wire_length(h, w)
    return (
        0 if w.gauge_awg is not None else 1,
        w.gauge_awg if w.gauge_awg is not None else 0,
        (w.colour or "~").lower(),
        0 if length is not None else 1,
        length if length is not None else 0.0,
        w.id,
    )


def _far(w: Wire, cid: str, pin: str) -> str:
    return f"{w.to_connector}.{w.to_pin}" if (w.from_connector, w.from_pin) == (cid, pin) else f"{w.from_connector}.{w.from_pin}"  # fmt: skip


def _attached(h: Harness) -> dict[tuple[str, str], list[Wire]]:
    out: dict[tuple[str, str], list[Wire]] = defaultdict(list)
    for w in h.wires:
        out[(w.from_connector, w.from_pin)].append(w)
        out[(w.to_connector, w.to_pin)].append(w)
    return out


@dataclass(frozen=True)
class Check:
    """One line of the *Start here* sheet: what is checked, how many are open, what to do."""

    label: str
    open: int
    blocking: bool
    todo: str

    @property
    def status(self) -> str:
        if self.open == 0:
            return OK
        return TO_FIX if self.blocking else NOTE


def readiness(project: Project, h: Harness) -> list[Check]:
    """What stands between this harness and a build, worked out from the model alone."""
    wires = list(h.wires)
    attached = _attached(h)
    used_pins = sum(
        1
        for c in h.connectors
        for p in c.pins
        if attached.get((c.id, p.id)) and p.termination == "unspecified"
    )
    no_contact = sum(
        1 for c in h.connectors for p in c.pins if attached.get((c.id, p.id)) and not p.contact_part
    )
    parts = {c.part_id for c in h.connectors} | {w.part_id for w in wires if w.part_id}
    not_approved = sum(
        1 for pid in parts if pid not in project.parts or project.parts[pid].approval != "approved"
    )
    return [
        Check("Wires without a length", sum(1 for w in wires if wire_length(h, w) is None), True,
              "Import lengths (import-lengths) or set the length of the wire in the Edit wire tab."),
        Check("Wires without a gauge", sum(1 for w in wires if w.gauge_awg is None), True,
              "Generate again with the engineering values filled in, or set the gauge of the wire."),
        Check("Wires without a colour", sum(1 for w in wires if not w.colour), True,
              "Set the colour of the signal (Edit > Wire colours) or of the wire."),
        Check("Wires without a part", sum(1 for w in wires if not w.part_id), True,
              "Choose the wire part in the Edit wire tab."),
        Check("Pins in use without a termination", used_pins, False,
              "Say crimp or solder for the pins in the Devices tab; the build sheet shows it."),
        Check("Wires without strip lengths",
              sum(1 for w in wires if w.strip_a_mm is None or w.strip_b_mm is None), False,
              "Enter the strip length at both ends in the Edit wire tab (or with import-build-data)."),
        Check("Pins in use without a contact part", no_contact, False,
              "Name the contact part of each pin (Devices tab, or Connectors tab of a harness made by hand)."),
        Check("Parts that are not approved", not_approved, False,
              "Approve the parts in the parts list (import-parts) when they have been reviewed."),
    ]  # fmt: skip


# ---- the sheets ------------------------------------------------------------------------------


def ready_text(answer: str, placeholders: list[str]) -> str:
    """The answer to *Ready to build*. A yes that rests on placeholder configuration says so
    (audit F-06); a no, or a project without placeholders, is the bare word as before."""
    if answer != "yes" or not placeholders:
        return answer
    return f"yes (values rest on placeholder configuration: {', '.join(placeholders)})"


def _start_sheet(wb: Any, project: Project, h: Harness, stamp: Stamp) -> None:
    ws = wb.create_sheet("Start here")
    lay = _Layout({})
    body: list[list[Any]] = []
    title = f"{h.id}  {h.label or h.name}"
    body.append([_cell(ws, title, bold=True, size=16, box=False, wrap=False)])
    body.append([
        _cell(ws, f"Revision {h.revision}   ·   status {h.status}   ·   {stamp.line}", box=False, wrap=False,
              colour=_GREY_TEXT)
    ])  # fmt: skip
    body.append([])
    checks = readiness(project, h)
    blocking = sum(1 for c in checks if c.status == TO_FIX)
    ready = "yes" if blocking == 0 else "no"
    placeholders = project.placeholder_configs()
    shown = ready_text(ready, placeholders)
    if shown != ready:
        lay.see(2, shown)  # a long answer gets a wider column, not a tall row
    body.append([
        _cell(ws, "Ready to build", bold=True, fill=_BAND), _cell(ws, shown, bold=True,
              fill=_GREEN if ready == "yes" else _RED),
        _cell(ws, "some points marked TO FIX below are still open" if ready == "no" else
              "nothing marked TO FIX is open", box=False, wrap=False, colour=_GREY_TEXT),
    ])  # fmt: skip
    body.append([])
    body.append(_header(ws, lay, CHECK_HEADER))
    for c in checks:
        fill = {OK: _GREEN, TO_FIX: _RED, NOTE: _AMBER}[c.status]
        todo = "" if c.status == OK else c.todo
        body.append([
            _cell(ws, c.label), _cell(ws, c.open, centre=True),
            _cell(ws, c.status, bold=True, fill=fill, centre=True), _cell(ws, todo),
        ])  # fmt: skip
        for n, v in enumerate((c.label, str(c.open), c.status, todo), start=1):
            lay.see(n, v)
    body.append([])
    body.append([_cell(ws, "From the design", bold=True, box=False, wrap=False)])
    for label, value in (
        ("Author", h.author), ("Checker", h.checker), ("Approver", h.approver),
        ("Released on", h.released_on),
    ):  # fmt: skip
        body.append([_cell(ws, label, fill=_BAND), _cell(ws, value or "not set")])
        lay.see(1, label)
    if placeholders:  # the files behind the answer above (audit F-06)
        files = ", ".join(f"config/{n}.json" for n in placeholders)
        body.append([_cell(ws, "Placeholder files", fill=_BAND), _cell(ws, files)])
        lay.see(1, "Placeholder files")
        lay.see(2, files)
    body.append([])
    body.append([_cell(ws, "How to use this workbook", bold=True, box=False, wrap=False)])
    for name, what in (
        ("Cut list", "cut the wires in this order; tick each when done"),
        ("Connectors", "one block per connector: the pins, the wire for each and the steps"),
        ("Parts", "what to take from stock and whether each part is approved"),
        ("Accessories", "sleeving, heat shrink, labels and backshells to fit; tick each when done"),
        ("Labels", "the labels to print and fit"),
        ("Tests", "continuity and insulation tests; write down the result"),
        ("About", "which version of the tool and which design this was made from"),
    ):  # fmt: skip
        body.append(
            [_cell(ws, name, bold=True, fill=_BAND), _cell(ws, what, box=False, wrap=False)]
        )
        lay.see(1, name)
    body.append([])
    body.append([_cell(ws, "Build sign-off (fill in by hand)", bold=True, box=False, wrap=False)])
    body.append(_header(ws, lay, ["Step", "Name", "Date", "Signature"]))
    for step in (
        "Cut and stripped", "Terminated", "Assembled and labelled", "Tested", "Released for delivery"
    ):  # fmt: skip
        body.append([_cell(ws, step, fill=_BAND), _cell(ws, ""), _cell(ws, ""), _cell(ws, "")])
        lay.see(1, step)
    lay.widths[1] = max(lay.widths.get(1, 0), 30.0)
    lay.widths[2] = max(lay.widths.get(2, 0), 18.0)
    lay.widths[3] = max(lay.widths.get(3, 0), 14.0)
    lay.widths[4] = min(60.0, max(lay.widths.get(4, 0), 24.0))
    _finish(ws, lay, stamp, body, tab="1F7A3A", header_row=None, one_page=True)


def _strip(v: float | None) -> float | str:
    """A strip length for a cell: the number, or *to fill in*."""
    return round(v, 9) + 0.0 if v is not None else TO_FILL


def _cut_sheet(wb: Any, project: Project, h: Harness, stamp: Stamp) -> None:
    ws = wb.create_sheet("Cut list")
    lay = _Layout({})
    body: list[list[Any]] = []
    shields = tables.shield_of(h)
    body.append(_header(ws, lay, CUT_HEADER))
    wires = sorted(h.wires, key=lambda w: _wire_sort_key(h, w))
    total = 0.0
    for w in wires:
        length = wire_length(h, w)
        if length is not None:
            total += length
        values = [
            "", w.id, w.colour or TO_FILL, w.gauge_awg if w.gauge_awg is not None else TO_FILL,
            w.part_id or TO_FILL, round(length, 9) + 0.0 if length is not None else TO_FILL,
            _strip(w.strip_a_mm), _strip(w.strip_b_mm),
            f"{w.from_connector}.{w.from_pin}", f"{w.to_connector}.{w.to_pin}", tables.wire_name(w),
            shields.get(w.id, ""), w.note,
        ]  # fmt: skip
        body.append([
            _cell(ws, "", centre=True), _cell(ws, w.id, bold=True), _wire_colour_cell(ws, w.colour),
            _cell(ws, values[3], fill=_AMBER if w.gauge_awg is None else None, centre=True),
            _cell(ws, values[4], fill=_AMBER if not w.part_id else None),
            _cell(ws, values[5], fill=_AMBER if length is None else None, centre=True),
            _cell(ws, values[6], fill=_AMBER if w.strip_a_mm is None else None, centre=True),
            _cell(ws, values[7], fill=_AMBER if w.strip_b_mm is None else None, centre=True),
            _cell(ws, values[8]), _cell(ws, values[9]), _cell(ws, values[10]),
            _cell(ws, values[11]), _cell(ws, values[12]),
        ])  # fmt: skip
        for n, v in enumerate(values, start=1):
            lay.see(n, _colour_text(w.colour) if n == 3 and w.colour else str(v))
    last = len(wires) + 1
    body.append([
        _cell(ws, "", fill=_BAND), _cell(ws, f"{len(wires)} wires", bold=True, fill=_BAND),
        *[_cell(ws, "", fill=_BAND) for _ in range(3)],
        _cell(ws, round(total, 9) + 0.0, bold=True, fill=_BAND, centre=True),
        *[_cell(ws, "", fill=_BAND) for _ in range(7)],
    ])  # fmt: skip
    ws.auto_filter.ref = f"A1:{_letter(len(CUT_HEADER))}{last}"
    if wires:
        _tick_rules(ws, f"A2:A{last}")
    lay.widths[1] = 10.0
    _finish(ws, lay, stamp, body, tab=_NAVY, freeze="C2")


def _pin_rows(
    ws: Any,
    project: Project,
    c: Connector,
    h: Harness,
    attached: dict[tuple[str, str], list[Wire]],
    lay: _Layout,
) -> list[list[Any]]:
    rows: list[list[Any]] = []
    for p in c.pins:
        wires = sorted(attached.get((c.id, p.id), []), key=lambda w: w.id)
        spare = not wires and not p.interface_id and not p.signal
        grey = _SPARE if not wires else None
        if wires:
            wire_ids = "; ".join(w.id for w in wires)
            goes = "; ".join(_far(w, c.id, p.id) for w in wires)
        else:
            wire_ids, goes = ("(spare)" if spare else "(no wire)"), ""
        gauge = "; ".join(str(w.gauge_awg) for w in wires if w.gauge_awg is not None)
        term = p.termination if p.termination != "unspecified" else (TO_FILL if wires else "")
        contact = p.contact_part or (TO_FILL if wires else "")
        part = project.parts.get(p.contact_part or "")
        tool = (part.tool or "") if part is not None else ""
        cells = [
            _cell(ws, p.id, bold=True, fill=grey), _cell(ws, p.signal or "", fill=grey),
            _cell(ws, p.function or "", fill=grey),
            _cell(ws, wire_ids, fill=grey, colour=_GREY_TEXT if not wires else None),
            _wire_colour_cell(ws, wires[0].colour) if len(wires) == 1 else
            _cell(ws, "; ".join(w.colour or "?" for w in wires), fill=grey),
            _cell(ws, gauge, fill=grey, centre=True), _cell(ws, goes, fill=grey),
            _cell(ws, term, fill=_AMBER if term == TO_FILL else grey),
            _cell(ws, p.contact_size or "", fill=grey),
            _cell(ws, contact, fill=_AMBER if contact == TO_FILL else grey),
            _cell(ws, tool, fill=grey),
            *[_cell(ws, "", fill=grey, centre=True) for _ in range(3)],
        ]  # fmt: skip
        if not wires:
            cells[4] = _cell(ws, "", fill=grey)
        rows.append(cells)
        shown_colour = (
            _colour_text(wires[0].colour)
            if len(wires) == 1
            else "; ".join(w.colour or "?" for w in wires)
        )
        for n, v in enumerate(
            (p.id, p.signal or "", p.function or "", wire_ids, shown_colour, gauge, goes, term, p.contact_size or "", contact, tool),
            start=1,
        ):  # fmt: skip
            lay.see(n, v)
    return rows


def _connector_sheet(wb: Any, project: Project, h: Harness, stamp: Stamp) -> None:
    from openpyxl.worksheet.pagebreak import Break

    ws = wb.create_sheet("Connectors")
    lay = _Layout({})
    body: list[list[Any]] = []
    attached = _attached(h)
    body.append(_header(ws, lay, CONNECTOR_HEADER))
    row, since_break, ticks = 1, 0, []
    for c in sorted(h.connectors, key=lambda x: x.id):
        label = tables.connector_label(project, c)
        head = [
            c.id, c.part_id, c.gender if c.gender != "unspecified" else "gender not set",
            f"mates with {c.mates_with}" if c.mates_with else "not plugged into a unit",
        ]  # fmt: skip
        if label:
            head.append(f"label: {label}")
        block = len(c.pins) + 2
        if since_break and since_break + block > 42:
            ws.row_breaks.append(Break(id=row))
            since_break = 0
        body.append([
            _cell(ws, "   ·   ".join(head), bold=True, fill=_NAVY, colour=_WHITE, size=11, wrap=False, box=False),
            *[_cell(ws, "", fill=_NAVY, box=False) for _ in range(len(CONNECTOR_HEADER) - 1)],
        ])  # fmt: skip
        row += 1
        rows = _pin_rows(ws, project, c, h, attached, lay)
        first = row + 1
        for cells in rows:
            body.append(cells)
            row += 1
        if rows:
            ticks.append((first, row))
        body.append([])
        row += 1
        since_break += block
    for first, last in ticks:
        _tick_rules(ws, f"{_letter(12)}{first}:{_letter(14)}{last}")
    lay.widths[1] = max(lay.widths.get(1, 0), 8.0)
    _finish(ws, lay, stamp, body, tab="B5651D", freeze="A2")


def _parts_sheet(wb: Any, project: Project, h: Harness, stamp: Stamp) -> None:
    ws = wb.create_sheet("Parts")
    lay = _Layout({})
    body: list[list[Any]] = []
    rows = tables.bom(project, [h])
    body.append(_header(ws, lay, PARTS_HEADER))
    for r in rows[1:]:
        approval = r[5]
        fill = (
            _GREEN if approval.startswith("approved")
            else _AMBER if approval.startswith("pending") else _RED
        )  # fmt: skip
        qty: object = r[6]
        with contextlib.suppress(ValueError):
            qty = float(r[6]) if r[7] == "m" else int(r[6])
        body.append([
            _cell(ws, r[0], bold=True), *[_cell(ws, v) for v in r[1:5]],
            _cell(ws, approval, fill=fill), _cell(ws, qty, centre=True), _cell(ws, r[7], centre=True),
            _cell(ws, r[8], centre=True), _cell(ws, r[9]),
        ])  # fmt: skip
        for n, v in enumerate(r, start=1):
            lay.see(n, v)
    ws.auto_filter.ref = f"A1:{_letter(len(PARTS_HEADER))}{len(rows)}"
    _finish(ws, lay, stamp, body, tab="7A4EB0", freeze="B2")


def _accessories_sheet(wb: Any, project: Project, h: Harness, stamp: Stamp) -> None:
    ws = wb.create_sheet("Accessories")
    lay = _Layout({})
    body: list[list[Any]] = []
    rows = tables.accessories(h)
    body.append(_header(ws, lay, ACCESSORY_HEADER))
    for r in rows[1:]:
        r = [
            *r[:3],
            r[3] or HARNESS_TARGET,
            *r[4:],
        ]  # the target of an accessory of the whole harness
        length = float(r[5]) if r[5] else ""
        body.append([
            _cell(ws, "", centre=True), _cell(ws, r[0], bold=True), _cell(ws, r[1]), _cell(ws, r[2]),
            _cell(ws, r[3]), _cell(ws, r[4]), _cell(ws, length, centre=True),
            _cell(ws, int(r[6]), centre=True), _cell(ws, r[7]),
        ])  # fmt: skip
        for n, v in enumerate(r, start=2):
            lay.see(n, v)
    if len(rows) == 1:  # no accessories: keep the header, say so in one line
        body.append([_cell(ws, NO_ACCESSORIES, box=False, wrap=False, colour=_GREY_TEXT)])
    ws.auto_filter.ref = f"A1:{_letter(len(ACCESSORY_HEADER))}{len(rows)}"
    if len(rows) > 1:
        _tick_rules(ws, f"A2:A{len(rows)}")
    for col, least in ((2, 12.0), (3, 12.0), (4, 18.0), (5, 14.0), (6, 28.0), (9, 28.0)):
        lay.widths[col] = max(
            lay.widths.get(col, 0.0), least
        )  # a sheet with no rows stays readable
    lay.widths[1] = 10.0
    _finish(ws, lay, stamp, body, tab="8E44AD", freeze="B2")


def _labels_sheet(wb: Any, project: Project, h: Harness, stamp: Stamp) -> None:
    ws = wb.create_sheet("Labels")
    lay = _Layout({})
    body: list[list[Any]] = []
    rows = tables.labels(project, h)
    body.append(_header(ws, lay, LABEL_HEADER))
    for r in rows[1:]:
        body.append([_cell(ws, "", centre=True), _cell(ws, r[0], bold=True), _cell(ws, r[1]),
                   _cell(ws, r[2]), _cell(ws, int(r[3]) if r[3].isdigit() else r[3], centre=True)])  # fmt: skip
        for n, v in enumerate(r, start=2):
            lay.see(n, v, cap=70.0)
    ws.auto_filter.ref = f"A1:{_letter(len(LABEL_HEADER))}{len(rows)}"
    if len(rows) > 1:
        _tick_rules(ws, f"A2:A{len(rows)}")
    lay.widths[1] = 10.0
    _finish(ws, lay, stamp, body, tab="2E8B8B", freeze="A2")


def _tests_sheet(wb: Any, project: Project, h: Harness, stamp: Stamp) -> None:
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.worksheet.datavalidation import DataValidation

    ws = wb.create_sheet("Tests")
    lay = _Layout({})
    body: list[list[Any]] = []
    rows = tables.tests(project, h)
    body.append(_header(ws, lay, TEST_HEADER))
    for r in rows[1:]:
        body.append([
            _cell(ws, r[0], bold=True), *[_cell(ws, v) for v in r[1:8]],
            *[_cell(ws, "", centre=True) for _ in range(4)],
        ])  # fmt: skip
        for n, v in enumerate(r, start=1):
            lay.see(n, v)
    last = len(rows)
    ws.auto_filter.ref = f"A1:{_letter(len(TEST_HEADER))}{last}"
    if last > 1:
        dv = DataValidation(type="list", formula1='"Pass,Fail"', allow_blank=True)
        ws.data_validations.append(dv)
        dv.add(f"J2:J{last}")
        ws.conditional_formatting.add(
            f"J2:J{last}", CellIsRule(operator="equal", formula=['"Pass"'], fill=_S["fill"](_GREEN))
        )
        ws.conditional_formatting.add(
            f"J2:J{last}", CellIsRule(operator="equal", formula=['"Fail"'], fill=_S["fill"](_RED))
        )
    for col in (9, 10, 11, 12):
        lay.widths[col] = 14.0
    _finish(ws, lay, stamp, body, tab="C0392B", freeze="B2")


def _about_sheet(wb: Any, stamp: Stamp) -> None:
    """The same two rows as the other workbooks, so the verifier reads the stamp the same way."""
    ws = wb.create_sheet("About")
    ws.append(["Generated by", stamp.line])
    ws.append(["Model hash", stamp.model_hash])


def build_workbook(project: Project, h: Harness, stamp: Stamp) -> bytes:
    """The build workbook of harness `h` (see the module text)."""
    wb = Workbook(write_only=True)
    wb.properties.creator = "harness-design-studio"
    wb.properties.title = f"{stamp.line} {h.id}"
    wb.properties.created = wb.properties.modified = _FIXED
    _start_sheet(wb, project, h, stamp)
    _cut_sheet(wb, project, h, stamp)
    _connector_sheet(wb, project, h, stamp)
    _parts_sheet(wb, project, h, stamp)
    _accessories_sheet(wb, project, h, stamp)
    _labels_sheet(wb, project, h, stamp)
    _tests_sheet(wb, project, h, stamp)
    _about_sheet(wb, stamp)
    raw = io.BytesIO()
    wb.save(raw)
    return normalise_xlsx(raw.getvalue())


def cut_rows(project: Project, h: Harness) -> Table:
    """The cut list as plain rows (what the sheet shows, in its order): the independent check and
    the tests compare the workbook with this."""
    shields = tables.shield_of(h)
    rows: Table = []
    for w in sorted(h.wires, key=lambda x: _wire_sort_key(h, x)):
        length = wire_length(h, w)
        rows.append([
            w.id, w.colour or "", "" if w.gauge_awg is None else str(w.gauge_awg), w.part_id or "",
            tables.data_num(length), tables.data_num(w.strip_a_mm), tables.data_num(w.strip_b_mm),
            f"{w.from_connector}.{w.from_pin}", f"{w.to_connector}.{w.to_pin}",
            tables.wire_name(w), shields.get(w.id, ""), w.note,
        ])  # fmt: skip
    return rows


__all__ = [
    "FILE_SUFFIX", "SHEETS", "Check", "build_workbook", "cut_rows", "readiness", "ready_text",
]  # fmt: skip
