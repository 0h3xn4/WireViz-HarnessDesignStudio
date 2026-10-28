"""Regression tests for the medium output findings of the October audit: D-3 (long pin IDs),
D-4 (reversible CSV guard), D-5/K-1 (title block fields), D-6 (PDF characters), D-7 (shield end),
D-8 (insertion order), D-10 (cell checks of the output verifier), D-13 (workbook errors), D-15,
D-16, D-18, G-7 (XLSX bytes with and without lxml) and I-4 (formula guard on every cell)."""

import hashlib
import io
import os
import re
import subprocess
import sys
import zipfile
from collections.abc import Callable
from pathlib import Path
from typing import Any

import pytest

from harness_design_studio.cli.main import main as cli_main
from harness_design_studio.core import configcheck, edit
from harness_design_studio.core.commands import Put, SetConfig, apply_ops
from harness_design_studio.core.generate.engine import generate_project
from harness_design_studio.core.io.saver import save_project
from harness_design_studio.core.model import (
    Connector,
    Harness,
    Pin,
    Project,
    Segment,
    ShieldGroup,
    Wire,
    evolve,
)
from harness_design_studio.core.model.config import ConfigFile
from harness_design_studio.core.outputs import drawing, exports, system
from harness_design_studio.core.outputs.build import OutputsError, build_outputs
from harness_design_studio.core.outputs.canvas import Sheet, Text, text_width
from harness_design_studio.core.outputs.stamp import Stamp, csv_bytes, parse_csv, stamp_of
from harness_design_studio.core.outputs.verify import _xlsx_sheets, verify_outputs
from harness_design_studio.core.samples import mini3, sat15, sat15_full
from tests.test_outputs import codes, edit_csv, project_and_files

STAMP = Stamp("0", "x" * 64)


def conn(cid: str, pins: list[str], signals: dict[str, str] | None = None) -> Connector:
    return Connector(
        id=cid,
        name=cid,
        role="cable",
        part_id="EX-DSUB-15-F",
        pins=[Pin(id=i, signal=(signals or {}).get(i)) for i in pins],
    )


def wire(wid: str, a: str, pa: str, b: str, pb: str) -> Wire:
    return Wire(id=wid, from_connector=a, from_pin=pa, to_connector=b, to_pin=pb)


def sheets_of(h: Harness, size: str = "A4", project: Project | None = None) -> list[Sheet]:
    return drawing.harness_sheets(project or Project(), h, STAMP, size)


# ---- D-3: pin IDs are shrunk to fit, never cut -----------------------------------------------------


def test_long_pin_ids_are_drawn_whole_and_do_not_touch_their_signal() -> None:
    pins_a = ["SHIELD_GND_1", "SHIELD_GND_2", "ABCDEFGHIJKLMNOP"]
    pins_b = ["P1", "P2", "P3"]
    h = Harness(
        id="H1",
        name="h",
        connectors=[conn("A", pins_a, {p: f"sig of {p}" for p in pins_a}), conn("B", pins_b)],
        wires=[wire(f"W{k}", "A", pins_a[k], "B", pins_b[k]) for k in range(3)],
    )
    items = [i for sh in sheets_of(h) for i in sh.items if isinstance(i, Text)]
    drawn = {i.ref: i for i in items if i.ref.startswith("pin/A/")}
    assert {r.removeprefix("pin/A/"): t.s for r, t in drawn.items()} == {p: p for p in pins_a}
    assert not any("~" in i.s for i in items if i.ref.startswith("pin/"))
    for pin in pins_a:  # the pin number and the signal text of one row never overlap
        num, sig = drawn[f"pin/A/{pin}"], next(i for i in items if i.ref == f"sig/A/{pin}")
        assert num.size >= drawing.ID_FLOOR
        assert sig.left + text_width(sig.s, sig.size) <= num.left + 0.01
        assert sig.left >= 0 and num.left + text_width(num.s, num.size) <= num.x + 0.01


def test_normal_pin_ids_keep_their_size_and_place() -> None:
    h = Harness(
        id="H1",
        name="h",
        connectors=[conn("A", ["1", "12"]), conn("B", ["1", "2"])],
        wires=[wire("W1", "A", "1", "B", "1"), wire("W2", "A", "12", "B", "2")],
    )
    pins = [i for sh in sheets_of(h) for i in sh.items if isinstance(i, Text) and i.bold]
    assert {i.size for i in pins if i.ref.startswith("pin/")} == {drawing.PIN_SIZE}


# ---- D-4: the CSV formula guard can be undone exactly ----------------------------------------------

GUARD_CELLS = [
    "=1+1", "+5V", "-12V", "@x", "'=1+1", "'@x", "'-x", "'+x", "''=x", "'", "'x", "x'", "-5",
    "+5", "-0.5e3", "-inf", "\t5", "\r5", "'\t5", "'-5", "- 5", "'", "", "a,b", '"-q"',
]  # fmt: skip


def test_every_guarded_cell_reads_back_as_written() -> None:
    table = [["h"], *[[c] for c in GUARD_CELLS if c != ""], [""]]
    assert parse_csv(csv_bytes(table, STAMP)) == table


@pytest.mark.parametrize("ch", ["=", "+", "-", "@", "\t", "\r"])
def test_each_formula_character_is_guarded_in_the_file(ch: str) -> None:
    raw = csv_bytes([["h"], [f"{ch}cmd"]], STAMP).decode()
    cells = [row[0] for row in __import__("csv").reader(io.StringIO(raw.split("\n", 1)[1]))]
    assert cells[1] == f"'{ch}cmd"


def test_plain_numbers_are_not_guarded() -> None:
    raw = csv_bytes([["h"], ["-5"], ["+2.5"], ["-0.5e3"]], STAMP).decode()
    assert "'" not in raw


@pytest.mark.parametrize("signal", ["'=1+1", "'@x", "'-x", "'+x", "'=x", "+5V", "-12V"])
def test_signals_that_look_like_formulas_export_and_verify(signal: str) -> None:
    p = mini3()
    h = p.harnesses["W001"]
    p.harnesses["W001"] = evolve(h, wires=[evolve(h.wires[0], signal=signal), *h.wires[1:]])
    files = build_outputs(p).files
    report = verify_outputs(p, files)
    assert report.ok, [i.render() for i in report.errors]
    row = parse_csv(files["harnesses/W001/wirelist.csv"])[1]
    assert row[1] == signal


# ---- D-5 / K-1: title block fields --------------------------------------------------------------------


def _with_fields(p: Project, fields: object) -> Project:
    old = p.config["titleblock"]
    p.config["titleblock"] = ConfigFile(
        name="titleblock", placeholder=old.placeholder, values={**old.values, "fields": fields}
    )
    return p


@pytest.mark.parametrize("field", list(configcheck.TITLEBLOCK_FIELDS))
def test_any_one_title_block_field_can_be_left_out(field: str) -> None:
    p = mini3()
    fields = [f for f in configcheck.TITLEBLOCK_FIELDS if f != field]
    _with_fields(p, fields)
    assert not [i for i in configcheck.validate(p) if i.object_id == "titleblock.fields"]
    files = build_outputs(p).files
    report = verify_outputs(p, files)
    assert report.ok, [i.render() for i in report.errors]
    svg = files["harnesses/W001/drawing_A3_s1.svg"].decode()
    assert (field.replace("_", " ").upper() in svg) is False


@pytest.mark.parametrize(
    ("bad", "words"),
    [
        (["bogus", "project"], "bogus"),
        (["project", "project"], "more than once"),
        ([1, 2], "text"),
        ("project", "list"),
    ],
)
def test_a_bad_title_block_list_is_a_named_configuration_error(bad: object, words: str) -> None:
    p = _with_fields(mini3(), bad)
    issues = [i for i in configcheck.validate(p) if i.object_id == "titleblock.fields"]
    assert len(issues) == 1 and issues[0].code == "config_invalid"
    assert "titleblock" in issues[0].message and words in issues[0].message
    # the drawing falls back to the default fields, so the export still verifies
    files = build_outputs(p).files
    assert verify_outputs(p, files).ok


def test_the_default_title_block_fields_are_valid_and_empty_means_default() -> None:
    p = mini3()
    assert configcheck.validate_titleblock(p) == []
    _with_fields(p, [])
    assert configcheck.validate_titleblock(p) == []
    assert configcheck.titleblock_fields(p) == list(configcheck.TITLEBLOCK_DEFAULT)


# ---- D-6: characters the PDF cannot show are named ----------------------------------------------------


def test_characters_outside_windows_1252_are_named_in_a_warning() -> None:
    p = mini3()
    h = p.harnesses["W001"]
    p.harnesses["W001"] = evolve(
        h, wires=[evolve(h.wires[0], signal="Rsense 10kΩ ±5%"), *h.wires[1:]]
    )
    pins = p.harnesses["W001"].wires[0]
    for c in p.harnesses["W001"].connectors:
        if c.id == pins.from_connector:
            c_new = evolve(
                c,
                pins=[
                    evolve(x, signal="Rsense 10kΩ") if x.id == pins.from_pin else x for x in c.pins
                ],
            )
            p.harnesses["W001"] = evolve(
                p.harnesses["W001"],
                connectors=[c_new if k.id == c.id else k for k in p.harnesses["W001"].connectors],
            )
    files = build_outputs(p).files
    report = verify_outputs(p, files)
    assert report.ok
    warn = [i for i in report.issues if i.code == "out_pdf_chars"]
    assert warn and "U+03A9" in warn[0].message and "±" not in warn[0].message
    assert "'?'" in warn[0].message
    # the SVG keeps the character
    assert "Ω" in files["harnesses/W001/drawing_A3_s1.svg"].decode()


def test_a_clean_project_has_no_pdf_character_warning() -> None:
    p, files = project_and_files("sat15")
    assert not [i for i in verify_outputs(p, files).issues if i.code == "out_pdf_chars"]


# ---- D-7: the end of a shield does not change between sheets ----------------------------------------


def _two_way_shield() -> Harness:
    out = [wire(f"W{k:03d}", "A", str(k), "B", str(k)) for k in range(1, 31)]
    back = [wire(f"W{k:03d}", "B", str(k - 30), "A", str(k - 30)) for k in range(31, 61)]
    ids = [w.id for w in (*out, *back)]
    return Harness(
        id="H1",
        name="h",
        connectors=[
            conn("A", [str(k) for k in range(1, 31)]),
            conn("B", [str(k) for k in range(1, 31)]),
        ],
        wires=[*out, *back],
        shields=[
            ShieldGroup(
                id="S1",
                kind="overall_shield",
                wire_ids=ids,
                end_a="backshell_360",
                end_b="floating",
            )
        ],
    )


def test_a_shield_end_stays_on_the_same_side_on_every_sheet() -> None:
    result = sheets_of(_two_way_shield(), "A4")
    assert len(result) >= 3
    sides = []
    for sh in result:
        words = [
            i for i in sh.items if isinstance(i, Text) and i.s in ("backshell 360", "floating")
        ]
        for t in words:
            if t.s == "backshell 360":
                sides.append("left" if t.anchor == "start" else "right")
    assert len(sides) >= 3 and len(set(sides)) == 1


def test_a_shield_end_a_is_the_from_end_of_its_lowest_wire_in_any_order() -> None:
    """End A of a shield is the From connector of its lowest wire ID: not the order of the wire
    list, not the order of `wire_ids`, not which connector the drawing puts on the left."""

    def end_a_side(h: Harness) -> str:
        (sheet,) = sheets_of(h, "A4")
        words = [i for i in sheet.items if isinstance(i, Text) and i.s in ("backshell 360",)]
        (word,) = words
        return "left" if word.anchor == "start" else "right"

    def harness(wires: list[Wire], ids: list[str]) -> Harness:
        return Harness(
            id="H1",
            name="h",
            connectors=[conn("A", ["1", "2"]), conn("B", ["1", "2"])],
            wires=wires,
            shields=[
                ShieldGroup(
                    id="S1",
                    kind="overall_shield",
                    wire_ids=ids,
                    end_a="backshell_360",
                    end_b="floating",
                )  # fmt: skip
            ],
        )

    w1, w2 = wire("W001", "A", "1", "B", "1"), wire("W002", "B", "2", "A", "2")
    sides = {
        end_a_side(harness(list(ws), list(ids)))
        for ws in ((w1, w2), (w2, w1))
        for ids in (["W001", "W002"], ["W002", "W001"])
    }
    assert len(sides) == 1
    flipped = harness([wire("W001", "B", "1", "A", "1"), w2], ["W001", "W002"])
    assert end_a_side(flipped) != sides.pop()  # the From end of the lowest wire decides


# ---- D-8: the order of the harnesses does not change the bytes ---------------------------------------


def test_outputs_do_not_depend_on_the_order_of_harnesses_or_wires() -> None:
    p, files = project_and_files("sat15")
    items = list(p.harnesses.items())
    p.harnesses.clear()
    p.harnesses.update(reversed(items))
    for hid, h in list(p.harnesses.items()):
        p.harnesses[hid] = evolve(h, wires=list(reversed(h.wires)))
    other = build_outputs(p).files
    assert sorted(other) == sorted(files)
    assert [k for k in files if other[k] != files[k]] == []


def _two_harnesses_on_the_same_box_pins(first: str, second: str, rename: bool = True) -> Project:
    """mini3 with a copy of W001 named W000 (own connector and wire IDs when `rename`), both
    plugged into the same box connectors; the harnesses are added in the given order."""
    p = mini3()
    h = p.harnesses["W001"]

    def new(text: str) -> str:
        return text.replace("W001", "W000") if rename else text

    copy = evolve(
        h,
        id="W000",
        connectors=[evolve(c, id=new(c.id)) for c in h.connectors],
        wires=[
            evolve(
                w,
                id=w.id.replace("W001", "W000"),
                from_connector=new(w.from_connector),
                to_connector=new(w.to_connector),
            )
            for w in h.wires
        ],
        segments=[],
        branch_points=[],
    )
    made = {"W001": h, "W000": copy}
    p.harnesses.clear()
    for hid in (first, second):
        p.harnesses[hid] = made[hid]
    return p


def test_two_wires_on_one_box_pin_are_listed_in_id_order() -> None:
    from harness_design_studio.core.outputs import tables

    a = tables.box_pinouts(_two_harnesses_on_the_same_box_pins("W001", "W000"))
    b = tables.box_pinouts(_two_harnesses_on_the_same_box_pins("W000", "W001"))
    assert a == b
    shared = [r[7] for r in a[1:] if ";" in r[7]]
    assert shared and all(c.split(";") == sorted(c.split(";")) for c in shared)


def test_mating_rows_with_one_cable_connector_ID_are_sorted_by_harness() -> None:
    from harness_design_studio.core.outputs import tables

    a = tables.mating_matrix(_two_harnesses_on_the_same_box_pins("W001", "W000", rename=False))
    b = tables.mating_matrix(_two_harnesses_on_the_same_box_pins("W000", "W001", rename=False))
    assert a == b
    assert [r[7] for r in a[1:] if r[7]][:2] == ["W000", "W001"]


def test_whole_output_set_is_the_same_for_either_order_of_harnesses() -> None:
    one = build_outputs(_two_harnesses_on_the_same_box_pins("W001", "W000")).files
    two = build_outputs(_two_harnesses_on_the_same_box_pins("W000", "W001")).files
    assert sorted(one) == sorted(two)
    assert [k for k in one if one[k] != two[k]] == []


def test_a_signal_name_keeps_double_spaces_in_the_labels() -> None:
    from harness_design_studio.core.outputs import tables

    p = mini3()
    h = p.harnesses["W001"]
    p.harnesses["W001"] = evolve(h, wires=[evolve(h.wires[0], signal="A  B"), *h.wires[1:]])
    rows = tables.labels(p, p.harnesses["W001"])
    assert any("A  B" in r[2] for r in rows)
    assert verify_outputs(p, build_outputs(p).files).ok


# ---- D-10: the output verifier checks the cells -----------------------------------------------------

ZZ = "ZZZ"
H = "harnesses/W010"

# (project, file, row, column, expected error code)
CELL_MUTATIONS = [
    ("sat15", f"{H}/pinouts.csv", 1, 1, "out_pinout"),  # role
    ("sat15", f"{H}/pinouts.csv", 1, 2, "out_pinout"),  # mates with
    ("sat15", f"{H}/pinouts.csv", 1, 3, "out_pinout"),  # part
    ("sat15", f"{H}/pinouts.csv", 1, 5, "out_pinout"),  # signal
    ("sat15", f"{H}/pinouts.csv", 1, 6, "out_pinout"),  # interface
    ("sat15", f"{H}/pinouts.csv", 1, 8, "out_pinout"),  # spare flag
    ("full", f"{H}/bom.csv", 1, 1, "out_bom"),  # category
    ("full", f"{H}/bom.csv", 1, 2, "out_bom"),  # manufacturer
    ("full", f"{H}/bom.csv", 1, 3, "out_bom"),  # part number
    ("full", f"{H}/bom.csv", 1, 4, "out_bom"),  # description
    ("full", f"{H}/bom.csv", 1, 5, "out_bom"),  # approval
    ("full", f"{H}/bom.csv", 1, 7, "out_bom"),  # unit
    ("full", f"{H}/bom.csv", 1, 8, "out_bom"),  # wires without length
    ("full", f"{H}/bom.csv", 1, 9, "out_bom"),  # used in
    ("full", "system/bom.csv", 1, 9, "out_bom"),  # used in, system
    ("full", "system/bom.csv", 1, 5, "out_bom"),  # approval, system
    ("full", f"{H}/tests.csv", 1, 2, "out_tests"),  # from
    ("full", f"{H}/tests.csv", 1, 4, "out_tests"),  # expected
    ("full", f"{H}/tests.csv", 1, 5, "out_tests"),  # continuity limit
    ("full", f"{H}/tests.csv", -1, 5, "out_tests"),  # isolation limit
    ("full", f"{H}/tests.csv", -1, 6, "out_tests"),  # test voltage
    ("full", f"{H}/tests.csv", -1, 7, "out_tests"),  # wires of a shield test
    ("full", f"{H}/labels.csv", 1, 2, "out_labels"),  # text
    ("full", f"{H}/labels.csv", -1, 2, "out_labels"),
    ("full", f"{H}/labels.csv", 1, 3, "out_labels"),  # quantity
    ("full", f"{H}/mass_length.csv", 1, 9, "out_mass"),  # missing data
    ("sat15", f"{H}/mass_length.csv", 1, 9, "out_mass"),
    ("full", "system/mass_length.csv", -1, 5, "out_mass"),  # total mass
    ("full", "system/mass_length.csv", -1, 9, "out_mass"),  # total note
    ("sat15", "system/mass_length.csv", 1, 9, "out_mass"),
    ("sat15", "system/mating_matrix.csv", 1, 1, "out_matrix"),  # unit
    ("sat15", "system/mating_matrix.csv", 1, 2, "out_matrix"),  # part
    ("sat15", "system/mating_matrix.csv", 1, 3, "out_matrix"),  # gender
    ("sat15", "system/mating_matrix.csv", 1, 5, "out_matrix"),  # cable part
    ("sat15", "system/mating_matrix.csv", 1, 6, "out_matrix"),  # cable gender
    ("sat15", "system/mating_matrix.csv", 1, 7, "out_matrix"),  # harness
    ("sat15", "system/traceability.csv", 1, 1, "out_trace"),  # name
    ("sat15", "system/traceability.csv", 1, 2, "out_trace"),  # type
    ("sat15", "system/traceability.csv", 1, 3, "out_trace"),  # from
    ("sat15", "system/traceability.csv", 1, 4, "out_trace"),  # to
    ("sat15", "system/traceability.csv", 1, 5, "out_trace"),  # harnesses
    ("sat15", "system/traceability.csv", 1, 6, "out_trace"),  # wires
    ("sat15", "system/box_pinouts.csv", 1, 3, "out_box_pinouts"),  # part
    ("sat15", "system/box_pinouts.csv", 1, 5, "out_box_pinouts"),  # signal
    ("sat15", "system/box_pinouts.csv", 1, 8, "out_box_pinouts"),  # spare
    ("sat15", "system/box_pinouts.csv", -1, 7, "out_box_pinouts"),  # wires
    ("sat15", "system/drc_findings.csv", 1, 2, "out_drc"),  # severity
    ("sat15", "system/drc_findings.csv", 1, 4, "out_drc"),  # statement
    ("sat15", "system/drc_findings.csv", 1, 5, "out_drc"),  # waived
]


_PRISTINE_CODES: dict[str, set[str]] = {}


def pristine_codes(name: str) -> set[str]:
    """The error codes the verifier gives for the unchanged output set of a reference project.
    The set is the same for every test, and the verifier needs about a second for it, so it is
    worked out once per process (the files and the project are never changed here)."""
    if name not in _PRISTINE_CODES:
        p, files = project_and_files(name)
        _PRISTINE_CODES[name] = codes(p, files)
    return set(_PRISTINE_CODES[name])


@pytest.mark.parametrize(("name", "rel", "row", "col", "code"), CELL_MUTATIONS)
def test_changing_one_cell_is_caught(name: str, rel: str, row: int, col: int, code: str) -> None:
    p, files = project_and_files(name)
    assert code not in pristine_codes(name)

    def change(t: list[list[str]]) -> None:
        t[row][col] = ZZ

    edit_csv(files, rel, change, p)
    assert code in codes(p, files), rel


def test_a_removed_box_pinout_row_is_caught() -> None:
    p, files = project_and_files("sat15")
    edit_csv(files, "system/box_pinouts.csv", lambda t: t.pop(2), p)
    assert "out_box_pinouts" in codes(p, files)


def test_a_wrong_wire_order_in_a_pinout_cell_is_caught() -> None:
    p = _two_harnesses_on_the_same_box_pins("W001", "W000")
    files = build_outputs(p).files
    assert "out_box_pinouts" not in codes(p, files)
    rows = parse_csv(files["system/box_pinouts.csv"])
    multi = next(r for r in rows[1:] if ";" in r[7])  # a box pin that carries two wires
    ids = multi[7].split(";")
    multi[7] = ";".join(reversed(ids))
    files["system/box_pinouts.csv"] = csv_bytes(rows, stamp_of(p))
    assert "out_box_pinouts" in codes(p, files)


@pytest.mark.parametrize(
    "mutate",
    [
        lambda b: b.replace(b'gauge: "', b'gauge: "9', 1),
        lambda b: re.sub(rb"length: [0-9.]+", b"length: 99", b, count=1),
        lambda b: re.sub(rb'(    - "[^"]+": \[)"([^"]+)"', rb'\1"ZZ"', b, count=1),
        lambda b: re.sub(rb"pincount: \d+", b"pincount: 99", b, count=1),
        lambda b: b.replace(b'pinlabels: ["', b'pinlabels: ["Z', 1),
        lambda b: b.replace(b"type:", b"type: x", 1),
    ],
)
def test_changed_yaml_values_are_caught(mutate: Callable[[bytes], bytes]) -> None:
    p, files = project_and_files("full")
    rel = next(k for k in sorted(files) if k.endswith("wireviz.yaml") and b'gauge: "' in files[k])
    new = mutate(files[rel])
    assert new != files[rel]
    files[rel] = new
    assert "out_yaml" in codes(p, files)


def test_a_swapped_pin_number_in_the_drawing_is_caught() -> None:
    p, files = project_and_files("full")
    svg = files[f"{H}/drawing_A3_s1.svg"].decode()
    pins = re.findall(r'(<text [^>]*data-ref="pin/([^"/]+)/([^"]+)"[^>]*>)([^<]*)(</text>)', svg)
    (a, ca, pa, ta, _), (b, cb, pb, tb, _) = next(
        (x, y) for x in pins for y in pins if x[1] == y[1] and x[2] != y[2]
    )
    swapped = (
        svg.replace(a + ta, a + "@@", 1).replace(b + tb, b + ta, 1).replace(a + "@@", a + tb, 1)
    )
    assert swapped != svg
    files[f"{H}/drawing_A3_s1.svg"] = swapped.encode()
    assert "out_drawing" in codes(p, files)


def test_a_changed_signal_text_in_the_drawing_is_caught() -> None:
    p, files = project_and_files("full")
    svg = files[f"{H}/drawing_A3_s1.svg"].decode()
    new = re.sub(r'(data-ref="sig/[^"]+"[^>]*>)([^<]+)', r"\1WRONG", svg, count=1)
    assert new != svg
    files[f"{H}/drawing_A3_s1.svg"] = new.encode()
    assert "out_drawing" in codes(p, files)


def test_a_pin_missing_from_the_drawing_is_caught() -> None:
    p, files = project_and_files("full")
    svg = files[f"{H}/drawing_A3_s1.svg"].decode()
    new = re.sub(r'<text [^>]*data-ref="pin/[^>]*>[^<]*</text>\n', "", svg, count=1)
    assert new != svg
    files[f"{H}/drawing_A3_s1.svg"] = new.encode()
    assert "out_drawing" in codes(p, files)


def test_pdf_text_other_than_wire_ids_is_checked() -> None:
    p, files = project_and_files("full")
    a3 = files[f"{H}/drawing_A3.pdf"]
    cid = p.harnesses["W010"].connectors[0].id.encode()
    files[f"{H}/drawing_A3.pdf"] = a3.replace(b"(" + cid + b")", b"(" + cid[:-1] + b"Z)", 1)
    assert files[f"{H}/drawing_A3.pdf"] != a3
    assert "out_drawing" in codes(p, files)
    p, files = project_and_files("full")
    a4 = files[f"{H}/drawing_A4.pdf"]
    files[f"{H}/drawing_A4.pdf"] = a4.replace(b"(" + cid + b")", b"(" + cid[:-1] + b"Z)")
    assert "out_drawing" in codes(p, files)
    p, files = project_and_files("full")
    pdf = files["system/block_diagram.pdf"]
    uid = next(iter(p.units)).encode()
    files["system/block_diagram.pdf"] = pdf.replace(b"(" + uid + b")", b"(" + uid + b"Z)", 1)
    assert "out_drawing" in codes(p, files)


@pytest.mark.parametrize(
    "mutate",
    [
        lambda s: s.replace('"parts": ', '"parts": 1 + ', 1),
        lambda s: re.sub(r'"unverified_parts": \d+', '"unverified_parts": 12345', s, count=1),
        lambda s: re.sub(r'"released_harnesses": \d+', '"released_harnesses": 77', s, count=1),
        lambda s: re.sub(r'"units": \d+', '"units": 77', s, count=1),
        lambda s: re.sub(r'"interfaces": \d+', '"interfaces": 77', s, count=1),
        lambda s: re.sub(r'"values_sha256_16": "[0-9a-f]+"', '"values_sha256_16": "0"', s, count=1),
        lambda s: re.sub(r'"tool_version": "[^"]*"', '"tool_version": "9.9"', s, count=1),
        lambda s: re.sub(r'"project": "[^"]*"', '"project": "other"', s, count=1),
    ],
)
def test_changed_provenance_values_are_caught(mutate: Callable[[str], str]) -> None:
    p, files = project_and_files("full")
    old = files["system/provenance.json"].decode()
    new = mutate(old)
    assert new != old
    files["system/provenance.json"] = new.encode()
    assert "out_provenance" in codes(p, files)


def test_a_workbook_sheet_that_differs_from_its_csv_is_caught() -> None:
    p, files = project_and_files("sat15")
    sheets = _xlsx_sheets(files["system/system.xlsx"])
    sheets["Box pinouts"][1][4] = ZZ
    files["system/system.xlsx"] = exports.xlsx_bytes(sheets, stamp_of(p))
    assert "out_xlsx" in codes(p, files)
    p, files = project_and_files("sat15")
    sheets = _xlsx_sheets(files[f"{H}/W010.xlsx"])
    sheets["Pinouts"][1][4] = ZZ
    files[f"{H}/W010.xlsx"] = exports.xlsx_bytes(
        {k: v for k, v in sheets.items() if k != "About"}, stamp_of(p)
    )
    assert "out_xlsx" in codes(p, files)


def test_the_json_export_must_match_the_parts_and_units() -> None:
    p, files = project_and_files("sat15")
    files["system/export.json"] = files["system/export.json"].replace(
        b'"parts": [', b'"parts": [{"id": "X"}, ', 1
    )
    assert "out_json" in codes(p, files)


# ---- D-13: a workbook that cannot be built is reported with its reason --------------------------------


def test_a_workbook_that_cannot_be_written_raises_a_plain_error(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from harness_design_studio.core.outputs import xlsx_env

    real_save = xlsx_env.Workbook.save

    def full(self: Any, *a: Any, **k: Any) -> None:
        real_save(self, *a, **k)  # let the library finish its work, then fail as a full disk does
        raise OSError(27, "File too large")

    monkeypatch.setattr(xlsx_env.Workbook, "save", full)
    with pytest.raises(OutputsError, match=r"Excel workbook could not be built \(File too large\)"):
        exports.xlsx_bytes({"T": [["a"]]}, STAMP)
    with pytest.raises(OutputsError):
        build_outputs(mini3())


def test_any_error_of_the_spreadsheet_library_becomes_an_outputs_error(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from harness_design_studio.core.errors import HarnessError
    from harness_design_studio.core.outputs import xlsx_env

    class Broken(Exception):
        pass

    real_save = xlsx_env.Workbook.save

    def boom(self: Any, *a: Any, **k: Any) -> None:
        real_save(self, *a, **k)
        raise Broken("IO_EFBIG")

    monkeypatch.setattr(xlsx_env.Workbook, "save", boom)
    with pytest.raises(HarnessError, match="IO_EFBIG"):
        exports.xlsx_bytes({"T": [["a"]]}, STAMP)


def test_the_export_command_says_why_the_workbook_failed(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
) -> None:
    from harness_design_studio.core.outputs import xlsx_env

    real_save = xlsx_env.Workbook.save

    def full(self: Any, *a: Any, **k: Any) -> None:
        real_save(self, *a, **k)
        raise OSError(27, "File too large")

    save_project(sat15_with_harnesses(), tmp_path / "p")
    monkeypatch.setattr(xlsx_env.Workbook, "save", full)
    assert cli_main(["export", str(tmp_path / "p")]) == 2
    err = capsys.readouterr().err
    assert "File too large" in err and "Traceback" not in err
    assert not (tmp_path / "p" / "outputs").exists()


def sat15_with_harnesses() -> Project:
    p = sat15()
    generate_project(p)
    return p


# ---- D-15, D-16 ------------------------------------------------------------------------------------


def test_the_harness_overview_draws_an_edge_for_every_pair_of_units() -> None:
    p = sat15()
    generate_project(p)
    boxes = sorted(p.connectors.values(), key=lambda c: c.id)
    three: list[Connector] = []
    for b in boxes:
        if b.unit_id and b.unit_id not in {x.unit_id for x in three}:
            three.append(b)
        if len(three) == 3:
            break
    h = p.harnesses["W001"]
    cons = [evolve(h.connectors[0], id=f"X{k}", mates_with=three[k].id) for k in range(3)]
    p.harnesses["W001"] = evolve(h, connectors=cons)
    edges = system._harness_edges(p)
    mine = {frozenset((e.a, e.b)) for e in edges if e.label == "W001"}
    units = {c.unit_id for c in three}
    assert len(units) == 3 and len(mine) == 3
    assert all(len(pair) == 2 for pair in mine)


def test_the_length_of_a_routing_segment_is_never_cut_in_the_sketch() -> None:
    n = 14
    conns = [
        Connector(id=f"C{k:02d}", name=f"C{k}", role="cable", part_id="EX-DSUB-9-M")
        for k in range(n)
    ]
    segs = [
        Segment(id=f"W001-L{k}", from_node=f"C{k:02d}", to_node=f"C{k + 1:02d}", length_m=12.5)
        for k in range(n - 1)
    ]
    h = Harness(id="H1", name="chain", connectors=conns, segments=segs, branch_points=[])
    sheet = Sheet(297.0, 210.0)
    drawing._draw_sketch(sheet, h, 20.0, 297.0)
    texts = [i for i in sheet.items if isinstance(i, Text)]
    assert sum(t.s == "12.5 m" for t in texts) == n - 1
    assert not any("~" in t.s for t in texts)


# ---- G-7: the workbook bytes do not depend on lxml ---------------------------------------------------

_SCRIPT = """
import hashlib, sys
if sys.argv[1] == "first":
    import openpyxl  # imported elsewhere first, as the data importer does
from harness_design_studio.core.outputs.exports import xlsx_bytes
from harness_design_studio.core.outputs.stamp import Stamp
data = xlsx_bytes({"Wire list": [["Wire", "x"], ["W1", "-2"], ["W2", "=1+1"], ["W3", ""]]}, Stamp("t", "a" * 64))
print(hashlib.sha256(data).hexdigest())
"""


@pytest.mark.parametrize("order", ["normal", "first"])
def test_workbook_bytes_are_the_same_with_and_without_lxml(order: str, tmp_path: Path) -> None:
    pytest.importorskip("lxml")
    results = set()
    for env_value in (None, "True", "False"):
        env = dict(os.environ)
        env.pop("OPENPYXL_LXML", None)
        if env_value is not None:
            env["OPENPYXL_LXML"] = env_value
        env["PYTHONHASHSEED"] = "0"
        done = subprocess.run(  # noqa: S603
            [sys.executable, "-P", "-c", _SCRIPT, order],
            env={**env, "PYTHONPATH": os.pathsep.join(sys.path)},
            capture_output=True,
            text=True,
            check=False,
            cwd=tmp_path,
        )
        assert done.returncode == 0, done.stderr
        results.add(done.stdout.strip())
    assert len(results) == 1


# ---- I-4: no cell of any output is a formula --------------------------------------------------------

TRIGGERS = ("=", "+", "-", "@", "\t", "\r")
FORMULA_NAMES = ["=1+1", "+cmd|' /C calc'!A0", "-2+3", "@SUM(1+1)", "\t=1+1", "'=x", "-x"]


_FORMULA_BUILT: dict[str, tuple[Project, dict[str, bytes]]] = {}


def _formula_built(name: str) -> tuple[Project, dict[str, bytes]]:
    """The project of `_formula_project(name)` and its output set, built once per process: two
    tests read the same set (building it takes about 3.5 s). The caller gets its own copy of the
    project and of the file table; the bytes are immutable."""
    if name not in _FORMULA_BUILT:
        p = _formula_project(name)
        _FORMULA_BUILT[name] = (p, build_outputs(p).files)
    p, files = _FORMULA_BUILT[name]
    return edit.clone_with(p, []), dict(files)


def _formula_project(name: str) -> Project:
    p = sat15_full()
    unit = next(iter(p.units.values()))
    h = p.harnesses["W010"]
    wires = [evolve(w, signal=name, colour=name) for w in h.wires[:2]]
    apply_ops(
        p,
        [
            Put("units", evolve(unit, name=name)),
            Put("harnesses", evolve(h, name=name, notes=name, wires=[*wires, *h.wires[2:]])),
        ],
    )
    return p


def _raw_csv_cells(data: bytes) -> list[str]:
    import csv

    lines = data.decode().splitlines(keepends=True)[1:]  # without the stamp line
    return [c for row in csv.reader(io.StringIO("".join(lines))) for c in row]


def _plain_number(cell: str) -> bool:
    return re.fullmatch(r"[+-]?(\d+\.?\d*|\.\d+)([eE][+-]?\d+)?", cell) is not None


@pytest.mark.parametrize("name", FORMULA_NAMES)
def test_no_cell_of_any_output_can_run_as_a_formula(name: str) -> None:
    _, files = _formula_built(name)
    seen_guarded = 0
    for rel, data in files.items():
        if rel.endswith(".csv"):
            for cell in _raw_csv_cells(data):
                if cell[:1] in TRIGGERS:
                    assert _plain_number(cell), (rel, cell)
                if cell.startswith("'") and cell[1:2] in TRIGGERS:
                    seen_guarded += 1
        if rel.endswith(".xlsx"):
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                for member in z.namelist():
                    if member.startswith("xl/worksheets/"):
                        xml = z.read(member)
                        assert b"<f>" not in xml and b"<f " not in xml, (rel, member)
    assert seen_guarded > 0 or name in ("'=x", "-x")


@pytest.mark.parametrize("name", ["=1+1", "-x", "+x", "@x", "\t=1"])
def test_workbook_cells_with_formula_characters_are_stored_as_text(name: str) -> None:
    from harness_design_studio.core.outputs.xlsx_env import load_workbook

    data = exports.xlsx_bytes({"T": [["h"], [name]]}, STAMP)
    ws = load_workbook(io.BytesIO(data))["T"]
    cell = ws.cell(row=2, column=1)
    assert cell.data_type != "f" and cell.value == name


@pytest.mark.slow
def test_every_csv_cell_of_the_hostile_names_reads_back() -> None:
    for name in FORMULA_NAMES:
        p, files = _formula_built(name)
        for rel, data in files.items():
            if rel.endswith(".csv"):
                table = parse_csv(data)
                assert csv_bytes(table, stamp_of(p)) == data, rel
        assert verify_outputs(p, files).ok


def test_the_set_of_output_hashes_is_stable_for_equal_input() -> None:
    p, files = project_and_files("mini3")
    again = build_outputs(p).files
    assert {k: hashlib.sha256(v).hexdigest() for k, v in again.items()} == {
        k: hashlib.sha256(v).hexdigest() for k, v in files.items()
    }
    assert apply_ops is not None and SetConfig is not None  # imported for the helpers above
